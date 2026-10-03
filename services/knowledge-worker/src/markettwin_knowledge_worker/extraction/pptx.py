"""Deterministic text extraction from PowerPoint presentations."""

from pathlib import Path
from typing import Protocol, cast
from zipfile import BadZipFile

import pptx
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.exc import PackageNotFoundError
from pptx.table import Table
from pptx.text.text import TextFrame

from markettwin_knowledge_worker.extraction.chunking import SourcePart, text_evidence
from markettwin_knowledge_worker.extraction.contracts import (
    ExtractionIssue,
    ExtractionResult,
)


class PptxExtractionError(RuntimeError):
    """Raised when a PowerPoint presentation cannot be extracted."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code


class _TextFrameShape(Protocol):
    @property
    def text_frame(self) -> TextFrame: ...


class _TableShape(Protocol):
    @property
    def table(self) -> Table: ...


class PptxExtractor:
    """Extract useful text from each PowerPoint slide."""

    name = "python-pptx"
    version = pptx.__version__

    def __init__(
        self,
        *,
        minimum_visual_slide_text_chars: int = 40,
    ) -> None:
        self._minimum_visual_slide_text_chars = (
            minimum_visual_slide_text_chars
        )

    def extract(self, path: Path) -> ExtractionResult:
        """Return one evidence unit per slide containing text."""

        try:
            presentation = Presentation(str(path))
        except (
            OSError,
            PackageNotFoundError,
            BadZipFile,
            ValueError,
            KeyError,
        ) as exc:
            raise PptxExtractionError(
                "parsing_failed",
                f"Unable to extract PowerPoint: {path.name}",
            ) from exc

        source_parts: list[SourcePart] = []
        issues: list[ExtractionIssue] = []

        slide_count = len(presentation.slides)

        for slide_number, slide in enumerate(
            presentation.slides,
            start=1,
        ):
            parts: list[str] = []
            has_visual_content = False

            for shape in slide.shapes:
                if shape.shape_type in {
                    MSO_SHAPE_TYPE.PICTURE,
                    MSO_SHAPE_TYPE.CHART,
                }:
                    has_visual_content = True

                if shape.has_text_frame:
                    text_shape = cast(_TextFrameShape, shape)
                    text = text_shape.text_frame.text.strip()

                    if text:
                        parts.append(text)

                if shape.has_table:
                    table_shape = cast(_TableShape, shape)

                    for row in table_shape.table.rows:
                        values = [
                            cell.text.strip()
                            for cell in row.cells
                        ]

                        if any(values):
                            parts.append(
                                " | ".join(values)
                            )

            text = "\n".join(parts).strip()

            if text:
                source_parts.append(
                    SourcePart(text=text, locator={"slide": slide_number})
                )

            if (
                has_visual_content
                and len(text)
                < self._minimum_visual_slide_text_chars
            ):
                issues.append(
                    ExtractionIssue(
                        code="needs_visual_fallback",
                        message=(
                            "Slide contains visual content "
                            "but little extractable text."
                        ),
                        source_locator={
                            "slide": slide_number,
                        },
                        requires_fallback=True,
                    )
                )

            elif not text:
                issues.append(
                    ExtractionIssue(
                        code="blank_slide",
                        message=(
                            "Slide contains no extractable "
                            "text or detected visual content."
                        ),
                        source_locator={
                            "slide": slide_number,
                        },
                    )
                )

        return ExtractionResult(
            source_path=path,
            units=text_evidence(
                source_parts, extractor_name=self.name, extractor_version=self.version
            ),
            issues=tuple(issues),
            source_item_count=slide_count,
            processed_item_count=slide_count,
        )
