"""Representative source formats preserve content, provenance, and coverage."""

import asyncio
import json
import shutil
from pathlib import Path
from unittest.mock import AsyncMock

import pytest
from markettwin_knowledge_worker.extraction import extract_source
from markettwin_knowledge_worker.extraction.csv_file import CsvExtractor
from markettwin_knowledge_worker.extraction.image import ImageExtractionError, ImageExtractor
from markettwin_knowledge_worker.extraction.json_file import JsonExtractor
from markettwin_knowledge_worker.extraction.pptx import PptxExtractor
from markettwin_knowledge_worker.extraction.text import TextExtractor
from markettwin_knowledge_worker.extraction.video import (
    VideoClip,
    VideoExtractor,
)
from markettwin_knowledge_worker.extraction.xlsx import XlsxExtractor
from openpyxl import Workbook
from pptx import Presentation


def _small_chunks(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setenv("KNOWLEDGE_SOURCE_CHUNK_MAX_CHARS", "80")
    monkeypatch.setenv("KNOWLEDGE_SKILL_BATCH_MAX_CHARS", "200")


def test_text_csv_and_json_use_bounded_model_ready_units(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _small_chunks(monkeypatch)
    text_path = tmp_path / "rules.md"
    text_path.write_text(
        "# Upload\nPDF and DOCX are supported.\nMaximum size is 10 MB.\n",
        encoding="utf-8",
    )
    csv_path = tmp_path / "limits.csv"
    csv_path.write_text("format,limit\nPDF,10 MB\nDOCX,10 MB\n", encoding="utf-8")
    json_path = tmp_path / "rules.json"
    json_path.write_text(
        json.dumps({"formats": ["PDF", "DOCX"], "rules": "x" * 180}), encoding="utf-8"
    )

    text_result = TextExtractor().extract(text_path)
    csv_result = CsvExtractor().extract(csv_path)
    json_result = JsonExtractor().extract(json_path)

    assert "10 MB" in "".join(unit.content_text or "" for unit in text_result.units)
    assert all(len(unit.content_text or "") <= 80 for unit in csv_result.units)
    assert all(unit.source_locator.get("row_start") for unit in csv_result.units)
    assert len(json_result.units) > 1
    assert all(unit.content_json is not None for unit in json_result.units)
    assert all(
        len(json.dumps(unit.content_json, separators=(",", ":"))) <= 80
        for unit in json_result.units
    )


def test_pptx_and_xlsx_chunk_without_losing_source_order(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _small_chunks(monkeypatch)
    pptx_path = tmp_path / "journey.pptx"
    presentation = Presentation()
    for title in ("Upload resume", "Download report"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        title_shape = slide.shapes.title
        assert title_shape is not None
        title_shape.text = title
    presentation.save(str(pptx_path))

    xlsx_path = tmp_path / "rules.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Upload rules"
    for row in (("Format", "Limit"), ("PDF", "10 MB"), ("DOCX", "10 MB")):
        sheet.append(row)
    workbook.save(xlsx_path)

    presentation_result = PptxExtractor().extract(pptx_path)
    workbook_result = XlsxExtractor().extract(xlsx_path)

    slide_text = "\n".join(unit.content_text or "" for unit in presentation_result.units)
    assert slide_text.index("Upload resume") < slide_text.index("Download report")
    assert presentation_result.coverage_complete
    assert all("Upload rules" in (unit.content_text or "") for unit in workbook_result.units)
    assert "10 MB" in "".join(unit.content_text or "" for unit in workbook_result.units)


@pytest.mark.parametrize("oversized_index", [0, 1, 2])
def test_json_array_ranges_match_payload_after_recursive_splitting(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch, oversized_index: int
) -> None:
    _small_chunks(monkeypatch)
    items: list[object] = ["first", "second", "third", "fourth"]
    items[oversized_index] = "x" * 180
    path = tmp_path / "array.json"
    path.write_text(json.dumps({"items": items}), encoding="utf-8")

    result = JsonExtractor().extract(path)

    split_content = ""
    grouped_indices: list[int] = []
    for unit in result.units:
        locator = str(unit.source_locator["json_path"])
        assert unit.content_json is not None
        payload = unit.content_json["value"]
        if "[chars:" in locator:
            assert locator.startswith(f"$.items[{oversized_index}][chars:")
            assert isinstance(payload, str)
            split_content += payload
        else:
            start, end = (int(index) for index in locator.removeprefix("$.items[")[:-1].split(":"))
            assert payload == items[start:end]
            grouped_indices.extend(range(start, end))
    assert split_content == items[oversized_index]
    assert grouped_indices == [index for index in range(len(items)) if index != oversized_index]


async def test_dispatcher_and_image_adapter_preserve_whole_media(
    tmp_path: Path,
) -> None:
    text_path = tmp_path / "source.TXT"
    text_path.write_text("Users can upload a resume.", encoding="utf-8")
    assert (await extract_source(text_path)).units

    image_path = tmp_path / "screen.png"
    image_path.write_bytes(b"image bytes")
    result = await ImageExtractor().extract(image_path)

    assert result.units[0].content_text is None
    assert result.units[0].content_json == {"media_type": "image/png"}
    assert result.units[0].source_locator == {"image": 1}
    assert result.units[0].evidence_type == "image"
    assert result.units[0].media_path == image_path
    assert result.coverage_complete


async def test_image_adapter_rejects_empty_images(tmp_path: Path) -> None:
    image_path = tmp_path / "empty.png"
    image_path.write_bytes(b"")
    with pytest.raises(ImageExtractionError, match="Image is empty"):
        await ImageExtractor().extract(image_path)


async def test_video_adapter_returns_timestamped_clips(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "demo.mp4"
    path.write_bytes(b"video")
    extractor = VideoExtractor()
    clips = (
        VideoClip(path, 0.0, 90.0),
        VideoClip(path, 85.0, 120.0),
    )
    def duration_seconds(_path: Path) -> float:
        return 120.0

    monkeypatch.setattr(extractor, "_duration_seconds", duration_seconds)
    monkeypatch.setattr(extractor, "_prepare_clips", AsyncMock(return_value=clips))

    result = await extractor.extract(path)

    assert result.source_item_count == 2
    assert result.processed_item_count == 2
    assert result.coverage_complete
    assert result.issues == ()
    assert [unit.source_locator for unit in result.units] == [
        {"start_seconds": 0.0, "end_seconds": 90.0},
        {"start_seconds": 85.0, "end_seconds": 120.0},
    ]
    assert [unit.media_path for unit in result.units] == [path, path]
    assert result.cleanup_paths
    shutil.rmtree(result.cleanup_paths[0], ignore_errors=True)


async def test_video_cancellation_cleans_prepared_clips(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    source = tmp_path / "source.mp4"
    source.write_bytes(b"original")
    clips_root = tmp_path / "clips"
    clips_root.mkdir()
    (clips_root / "partial.mp4").write_bytes(b"partial")
    extractor = VideoExtractor()

    def duration_seconds(_path: Path) -> float:
        return 120.0

    def temporary_directory(*, prefix: str) -> str:
        assert prefix == "markettwin-video-"
        return str(clips_root)

    monkeypatch.setattr(extractor, "_duration_seconds", duration_seconds)
    monkeypatch.setattr(
        "markettwin_knowledge_worker.extraction.video.mkdtemp", temporary_directory
    )
    monkeypatch.setattr(
        extractor, "_prepare_clips", AsyncMock(side_effect=asyncio.CancelledError)
    )
    with pytest.raises(asyncio.CancelledError):
        await extractor.extract(source)
    assert not clips_root.exists()
    assert source.read_bytes() == b"original"
