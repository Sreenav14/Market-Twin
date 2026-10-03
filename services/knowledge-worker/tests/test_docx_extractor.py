"""DOCX source order and the exact NumPy efficiency regression."""

from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from markettwin_knowledge_worker.extraction.docx import DocxExtractor


def test_docx_keeps_tables_between_paragraphs(tmp_path: Path) -> None:
    path = tmp_path / "ordered.docx"
    document = Document()
    document.add_heading("Resume upload", level=1)
    document.add_paragraph("Before the pricing table")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Plan"
    table.cell(0, 1).text = "Uploads"
    table.cell(1, 0).text = "Pro"
    table.cell(1, 1).text = "Yes"
    document.add_paragraph("After the pricing table")
    document.save(str(path))
    result = DocxExtractor().extract(path)
    assert len(result.units) == 1
    text = result.units[0].content_text or ""
    assert text.index("Before") < text.index("[TABLE]") < text.index("After")
    assert "# Resume upload" in text and "Pro | Yes" in text
    assert result.units[0].source_locator == {
        "paragraph_start": 1,
        "paragraph_end": 3,
        "table_start": 1,
        "table_end": 1,
    }
    assert result.source_item_count == result.processed_item_count == 4


def test_numpy_docx_retains_every_meaningful_part() -> None:
    path = Path(__file__).resolve().parents[3] / "numpy.docx"
    document = Document(str(path))
    result = DocxExtractor().extract(path)
    assert 1 <= len(result.units) <= 3
    assert not result.issues
    text = "\n".join(unit.content_text or "" for unit in result.units)
    for item in document.iter_inner_content():
        if isinstance(item, Paragraph) and item.text.strip():
            assert item.text.strip() in text
        elif isinstance(item, Table):
            for row in item.rows:
                assert " | ".join(cell.text.strip() for cell in row.cells) in text
